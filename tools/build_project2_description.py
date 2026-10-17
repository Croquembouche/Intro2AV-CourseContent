#!/usr/bin/env python3
"""Apply Project 2 wording to Project 1's retained document styles and geometry."""
from copy import deepcopy
import argparse
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'output/docx/Project1_Description.docx'
W = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def text_of(element):
    return ''.join(element.itertext()) if element is None else ''.join(element.xpath('.//w:t/text()', namespaces=W))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project-number', choices=['2', '2.2'], default='2')
    number = parser.parse_args().project_number
    slug = number.replace('.', '_')
    SOURCE = ROOT / f'2026/Data/Project{number}/Project{slug}_Description.md'
    OUTPUT = ROOT / f'output/docx/Project{slug}_Description.docx'
    QA = ROOT / f'tmp/docx/project{slug}'
    title = SOURCE.read_text().splitlines()[0].removeprefix('# ').replace(':', '')
    submission_contract = (
        "Each group submits a constructed map and one one-page technical report; each member separately submits an individual contribution summary report."
        if number == '2.2' else
        "The only deliverables are the constructed map and a one-page technical report."
    )
    QA.mkdir(parents=True, exist_ok=True)
    original_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    reference = Document(REFERENCE)
    prototypes = {}
    for paragraph in reference.paragraphs:
        style = paragraph.style.name
        if style == 'Normal' and paragraph.text.startswith('You are a perception'):
            prototypes['Normal'] = deepcopy(paragraph._p)
        elif style in ['Title', 'Heading 1', 'Heading 2', 'List Bullet']:
            prototypes.setdefault(style, deepcopy(paragraph._p))
        if paragraph.text in ['Required ROS 2 output', 'What to submit']:
            prototypes[paragraph.text] = deepcopy(paragraph._p)
    section = reference.sections[0]
    styles = {}
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
        style = reference.styles[name]
        styles[name] = {'font': style.font.name, 'size_pt': style.font.size.pt if style.font.size else None,
                        'xml': etree.tostring(style.element).decode()}
    with ZipFile(REFERENCE) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    editable = {'word/document.xml', 'word/footer1.xml', 'docProps/core.xml', 'word/settings.xml'}
    inventory = {name: {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(),
                        'role': 'editable' if name in editable else 'preserve-only'} for name, data in parts.items()}
    contract = f'''# Project {number} document execution contract

Reference: `{REFERENCE}`. SHA-256: `{original_hash}`. Reference has one section and three rendered pages, inspected in `tmp/docx/project1/render/`. Style and section audits are recorded under this directory.

Retain Letter portrait geometry, 0.75-inch side margins, 0.65-inch top/bottom margins, and 0.3-inch header/footer distances. Retain the reference's Calibri typography, black headings, 11-point Normal text, 22-point title, 14-point Heading 1 and 12-point Heading 2. Retain reference styles, bullet numbering, paragraph rhythm, and right-aligned page footer. Code spans use 10-point Consolas as in Project 1. Reference paragraphs, table cells, footer and PAGE field were inspected; there are no images or controls to fill.

Rewrite body content using Title, Normal, Heading 1, Heading 2 and List Bullet prototypes. Start Mapping requirements on the second page using the source's output-section heading pattern; let What to submit flow naturally. Remove the Project 1 output field table. {submission_contract} Omit the assessment section and any other submission requirements. Clone source paragraph patterns when Project {number} requires more paragraphs. Replace the course project number in the footer and metadata title, and request PAGE-field updates on open. All other package parts and relationships are preserve-only. Detailed package and typography evidence is in `template-evidence.json`.

The source Markdown is authoritative for Project {number} wording and due date. No Project 1 task, deadline or output interface may remain. Every source paragraph must be represented once. Verify unchanged package parts by SHA-256, retain reference bytes, render all final pages, and inspect typography, wrapping, footer placement and pagination.
'''
    (QA / 'artifact.md').write_text(contract)
    (QA / 'template-evidence.json').write_text(json.dumps({'reference_sha256': original_hash,
        'section_count': len(reference.sections), 'page_width_in': section.page_width.inches,
        'page_height_in': section.page_height.inches, 'styles': styles, 'parts': inventory}, indent=2))
    document = etree.fromstring(parts['word/document.xml'])
    body = document.find('w:body', W)
    sect = deepcopy(body.find('w:sectPr', W))
    for child in list(body):
        body.remove(child)
    expected = []
    for line in SOURCE.read_text().splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith('# '):
            role, text = 'Title', line[2:].replace(':', '')
        elif line.startswith('## '):
            text = line[3:]
            role = 'Required ROS 2 output' if text == 'Mapping requirements' else 'Heading 1'
        elif line.startswith('### '):
            role, text = 'Heading 2', re.sub(r'^(\d+)\. ', r'\1 ', line[4:])
            text = text.replace('One-page', 'One page')
        elif line.startswith('- '):
            role, text = 'List Bullet', line[2:]
        else:
            role, text = 'Normal', line
        paragraph = deepcopy(prototypes[role])
        for child in list(paragraph):
            if child.tag != qn('w:pPr'):
                paragraph.remove(child)
        plain = ''
        for segment in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
            if not segment:
                continue
            bold, code = segment.startswith('**'), segment.startswith('`')
            value = segment[2:-2].replace('`', '') if bold else segment[1:-1] if code else segment
            plain += value
            run = OxmlElement('w:r')
            if bold or code:
                props = OxmlElement('w:rPr')
                if bold:
                    props.append(OxmlElement('w:b'))
                if code:
                    fonts = OxmlElement('w:rFonts')
                    fonts.set(qn('w:ascii'), 'Consolas')
                    fonts.set(qn('w:hAnsi'), 'Consolas')
                    props.append(fonts)
                    size = OxmlElement('w:sz')
                    size.set(qn('w:val'), '20')
                    props.append(size)
                run.append(props)
            node = OxmlElement('w:t')
            node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            node.text = value
            run.append(node)
            paragraph.append(run)
        expected.append(plain)
        body.append(paragraph)
    body.append(sect)
    parts['word/document.xml'] = etree.tostring(document, xml_declaration=True, encoding='UTF-8', standalone=True)
    footer = etree.fromstring(parts['word/footer1.xml'])
    for node in footer.findall('.//w:t', W):
        if node.text:
            node.text = node.text.replace('Project 1', f'Project {number}')
    parts['word/footer1.xml'] = etree.tostring(footer, xml_declaration=True, encoding='UTF-8', standalone=True)
    core = etree.fromstring(parts['docProps/core.xml'])
    core.find('{http://purl.org/dc/elements/1.1/}title').text = title
    parts['docProps/core.xml'] = etree.tostring(core, xml_declaration=True, encoding='UTF-8', standalone=True)
    settings = etree.fromstring(parts['word/settings.xml'])
    update = settings.find('w:updateFields', W)
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')
    parts['word/settings.xml'] = etree.tostring(settings, xml_declaration=True, encoding='UTF-8', standalone=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(REFERENCE) as original, ZipFile(OUTPUT, 'w') as output:
        for entry in original.infolist():
            output.writestr(entry, parts[entry.filename])
    final = Document(OUTPUT)
    actual = [paragraph.text for paragraph in final.paragraphs]
    assert expected == actual, 'Source paragraph coverage mismatch'
    assert not any('Project 1' in paragraph or 'September 30' in paragraph for paragraph in actual)
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == original_hash
    with ZipFile(OUTPUT) as archive:
        changed = [name for name, row in inventory.items() if hashlib.sha256(archive.read(name)).hexdigest() != row['sha256']]
    assert set(changed) <= editable
    report = {'status': 'passed', 'paragraphs': len(actual), 'source_paragraphs_preserved': True,
              'reference_unchanged': True, 'preserve_only_parts_unchanged': True,
              'changed_package_parts': changed, 'output': str(OUTPUT),
              'source_words': len(SOURCE.read_text().split())}
    (QA / 'document_validation.json').write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
